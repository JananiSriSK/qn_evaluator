import os
import io
import logging
import math
from groq import Groq
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm, mm, inch
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn as dqn
from docx.oxml import OxmlElement

logger = logging.getLogger(__name__)

BT_ORDER = ["BT1", "BT2", "BT3", "BT4", "BT5", "BT6"]
BL_MAP   = {"BT1": "L1", "BT2": "L2", "BT3": "L3", "BT4": "L4", "BT5": "L5", "BT6": "L6"}

BT_VERBS = {
    "BT1": "Define / State / List / Recall",
    "BT2": "Explain / Describe / Summarize",
    "BT3": "Apply / Illustrate / Solve / Demonstrate",
    "BT4": "Analyze / Compare / Differentiate / Examine",
    "BT5": "Evaluate / Justify / Critique / Assess",
    "BT6": "Design / Create / Construct / Formulate",
}

BLOOM_LEGEND = (
    "BL – Bloom's Taxonomy Levels  "
    "(L1–Remembering, L2–Understanding, L3–Applying, "
    "L4–Analysing, L5–Evaluating, L6–Creating)"
)


class PaperGeneratorService:
    def __init__(self, groq_api_key=None):
        api_key = groq_api_key or os.getenv("GROQ_API_KEY")
        if not api_key:
            raise ValueError("GROQ_API_KEY not found")
        self.groq = Groq(api_key=api_key)

    # ── Slot builder ──────────────────────────────────────────────────────────

    def build_slots(self, syllabus: dict, parts: list) -> dict:
        units    = syllabus.get("units", [])
        cos      = syllabus.get("course_outcomes", {})
        co_keys  = sorted(cos.keys())
        n_units  = len(units)
        n_cos    = len(co_keys)
        n_parts  = len(parts)
        if not units:
            raise ValueError("Syllabus has no units")

        bt_ranges   = self._assign_bt_ranges(n_parts)
        unit_topics = {u["unit_number"]: [t["topic_name"] for t in u.get("topics", [])] for u in units}

        result_parts = []
        global_q = 0
        for part_idx, part_cfg in enumerate(parts):
            marks      = part_cfg["marks_per_question"]
            count      = part_cfg["question_count"]
            bt_pool    = bt_ranges[part_idx]
            unit_slots = self._distribute(count, n_units)

            slots      = []
            q_in_part  = 0
            for u_idx, unit in enumerate(units):
                for _ in range(unit_slots[u_idx]):
                    bt     = bt_pool[q_in_part % len(bt_pool)]
                    co     = co_keys[global_q % n_cos] if n_cos else ""
                    topics = unit_topics.get(unit["unit_number"], [])
                    topic  = topics[q_in_part % len(topics)] if topics else ""
                    slots.append({
                        "q_no": global_q + 1,
                        "part": part_cfg["part"],
                        "marks": marks,
                        "unit": unit["unit_number"],
                        "unit_title": unit["unit_title"],
                        "topic": topic,
                        "bloom": bt,
                        "co": co,
                    })
                    q_in_part += 1
                    global_q  += 1

            result_parts.append({
                "part": part_cfg["part"],
                "marks_per_question": marks,
                "slots": slots,
            })

        return {
            "course_name": syllabus.get("course_name", ""),
            "units": [{"number": u["unit_number"], "title": u["unit_title"],
                       "topics": [t["topic_name"] for t in u.get("topics", [])]} for u in units],
            "co_keys": co_keys,
            "parts": result_parts,
        }

    # ── Paper generator ───────────────────────────────────────────────────────

    def generate_paper(self, syllabus: dict, slots_by_part: list) -> dict:
        co_keys     = sorted(syllabus.get("course_outcomes", {}).keys())
        subtopic_map = {}
        for unit in syllabus.get("units", []):
            for t in unit.get("topics", []):
                subtopic_map[t["topic_name"]] = t.get("enriched_subtopics_from_books", [])

        paper_parts = []
        for part_cfg in slots_by_part:
            questions = []
            for slot in part_cfg["slots"]:
                subtopics = subtopic_map.get(slot["topic"], [])
                q_text    = self._generate_question(slot["topic"], subtopics, slot["bloom"], slot["marks"], slot["unit_title"])
                questions.append({**slot, "question": q_text})
            paper_parts.append({
                "part": part_cfg["part"],
                "marks_per_question": part_cfg["marks_per_question"],
                "question_count": len(questions),
                "questions": questions,
            })

        return {
            "course_name": syllabus.get("course_name", ""),
            "parts": paper_parts,
            "summary": self._build_summary(paper_parts, co_keys),
        }

    def regenerate_question(self, syllabus: dict, slot: dict) -> str:
        subtopic_map = {}
        for unit in syllabus.get("units", []):
            for t in unit.get("topics", []):
                subtopic_map[t["topic_name"]] = t.get("enriched_subtopics_from_books", [])
        subtopics = subtopic_map.get(slot["topic"], [])
        return self._generate_question(slot["topic"], subtopics, slot["bloom"], slot["marks"], slot["unit_title"])

    # ── Export — Anna University template ────────────────────────────────────

    def export_pdf(self, paper: dict, meta: dict = None) -> bytes:
        meta     = meta or {}
        usable_w = A4[0] - 40 * mm
        buf      = io.BytesIO()
        doc      = SimpleDocTemplate(buf, pagesize=A4,
                                     leftMargin=20*mm, rightMargin=20*mm,
                                     topMargin=15*mm, bottomMargin=15*mm)
        styles   = getSampleStyleSheet()

        def s(name, **kw):
            return ParagraphStyle(name, parent=styles["Normal"], **kw)

        cb  = s("cb",  fontSize=11, fontName="Helvetica-Bold", alignment=1, spaceAfter=1)
        cr  = s("cr",  fontSize=10, alignment=1, spaceAfter=1)
        sm  = s("sm",  fontSize=8)
        sb  = s("sb",  fontSize=8,  fontName="Helvetica-Bold")
        ps  = s("ps",  fontSize=11, fontName="Helvetica-Bold", spaceBefore=8, spaceAfter=4)
        qns = s("qns", fontSize=10, leading=14, spaceAfter=2)
        bls = s("bls", fontSize=7,  spaceAfter=6)

        story = []

        # University header
        story.append(Paragraph("ANNA UNIVERSITY, CHENNAI", cb))
        story.append(Paragraph("DEPARTMENT OF INFORMATION SCIENCE AND TECHNOLOGY", cb))
        story.append(Spacer(1, 3))
        story.append(Paragraph(meta.get("assessment", "Question Paper"), cr))
        story.append(Spacer(1, 6))

        # Info grid
        course_title = meta.get("course_title", "") or paper.get("course_name", "")
        info = Table([
            [Paragraph(f"<b>Programme:</b> {meta.get('programme', '')}", sm),
             Paragraph(f"<b>Year / SEM:</b> {meta.get('semester', '')}", sm)],
            [Paragraph(f"<b>Max. Marks:</b> {paper['summary']['total_marks']}", sm),
             Paragraph(f"<b>Regulation:</b> {meta.get('regulation', '')}", sm)],
            [Paragraph(f"<b>Date of Exam:</b> {meta.get('date', '')}", sm),
             Paragraph(f"<b>Duration:</b> {meta.get('duration', '')}", sm)],
        ], colWidths=[usable_w * 0.5, usable_w * 0.5])
        info.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))
        story.append(info)
        story.append(Spacer(1, 5))
        code = meta.get('course_code', '').strip()
        title = course_title.strip()
        code_title = f"{code} – {title}" if code else title
        story.append(Paragraph(f"<b>Course Code and Title:</b> {code_title}", cr))
        story.append(Spacer(1, 6))

        # CO table
        cos = meta.get("cos", [])
        if cos:
            co_data = [[Paragraph("<b>CO</b>", sb), Paragraph("<b>Description</b>", sb)]]
            for co in cos:
                co_data.append([Paragraph(co.get("id", ""), sm), Paragraph(co.get("text", ""), sm)])
            co_tbl = Table(co_data, colWidths=[usable_w * 0.10, usable_w * 0.90])
            co_tbl.setStyle(TableStyle([
                ("GRID",          (0, 0), (-1, -1), 0.5, colors.black),
                ("BACKGROUND",    (0, 0), (-1,  0), colors.HexColor("#4a5568")),
                ("TEXTCOLOR",     (0, 0), (-1,  0), colors.white),
                ("ALIGN",         (0, 0), (0,  -1), "CENTER"),
                ("VALIGN",        (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING",    (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(co_tbl)
            story.append(Spacer(1, 4))

        story.append(Paragraph(BLOOM_LEGEND, bls))

        # Questions by part
        import re as _re
        for part in paper["parts"]:
            story.append(Paragraph(f"Part {part['part']}  ({part['marks_per_question']} marks each)", ps))
            headers = ["Q.No", "Question", "CO", "BL", "Marks"]
            col_w   = [usable_w*0.06, usable_w*0.62, usable_w*0.10, usable_w*0.08, usable_w*0.14]
            rows    = [[Paragraph(f"<b>{h}</b>", sb) for h in headers]]
            for q in part["questions"]:
                raw_co = q.get("co", "")
                m = _re.match(r'(CO\d+)', str(raw_co), _re.IGNORECASE)
                co_label = m.group(1).upper() if m else raw_co
                rows.append([
                    Paragraph(str(q["q_no"]), sm),
                    Paragraph(q["question"], qns),
                    Paragraph(co_label, sm),
                    Paragraph(BL_MAP.get(q["bloom"], q["bloom"]), sm),
                    Paragraph(str(q["marks"]), sm),
                ])
            tbl = Table(rows, colWidths=col_w)
            tbl.setStyle(TableStyle([
                ("BACKGROUND",    (0, 0), (-1,  0), colors.HexColor("#4a5568")),
                ("TEXTCOLOR",     (0, 0), (-1,  0), colors.white),
                ("GRID",          (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN",        (0, 0), (-1, -1), "TOP"),
                ("ALIGN",         (0, 0), (0,  -1), "CENTER"),
                ("ALIGN",         (2, 0), (4,  -1), "CENTER"),
                ("TOPPADDING",    (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS",(0, 1), (-1, -1), [colors.white, colors.HexColor("#f7fafc")]),
            ]))
            story.append(tbl)
            story.append(Spacer(1, 0.2 * inch))

        doc.build(story)
        return buf.getvalue()

    def export_docx(self, paper: dict, meta: dict = None) -> bytes:
        meta         = meta or {}
        course_title = meta.get("course_title", "") or paper.get("course_name", "")

        doc = Document()
        for sec in doc.sections:
            sec.top_margin = sec.bottom_margin = Inches(0.6)
            sec.left_margin = sec.right_margin = Inches(0.8)

        def _p(text, bold=False, size=11, align=WD_ALIGN_PARAGRAPH.CENTER):
            p = doc.add_paragraph()
            p.alignment = align
            p.paragraph_format.space_after = p.paragraph_format.space_before = Pt(1)
            r = p.add_run(text)
            r.bold = bold
            r.font.size = Pt(size)

        def _shd(cell, fill):
            shd = OxmlElement("w:shd")
            shd.set(dqn("w:val"), "clear")
            shd.set(dqn("w:color"), "auto")
            shd.set(dqn("w:fill"), fill)
            cell._tc.get_or_add_tcPr().append(shd)

        def _no_borders(tbl):
            for row in tbl.rows:
                for cell in row.cells:
                    tcPr    = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement("w:tcBorders")
                    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                        b = OxmlElement(f"w:{side}")
                        b.set(dqn("w:val"), "none")
                        borders.append(b)
                    tcPr.append(borders)

        # University header
        _p("ANNA UNIVERSITY, CHENNAI", bold=True, size=12)
        _p("DEPARTMENT OF INFORMATION SCIENCE AND TECHNOLOGY", bold=True, size=11)
        _p(meta.get("assessment", "Question Paper"), size=10)

        # Info grid
        info = doc.add_table(rows=3, cols=2)
        info.alignment = WD_TABLE_ALIGNMENT.CENTER
        rows_data = [
            (f"Programme: {meta.get('programme', '')}",    f"Year / SEM: {meta.get('semester', '')}"),
            (f"Max. Marks: {paper['summary']['total_marks']}", f"Regulation: {meta.get('regulation', '')}"),
            (f"Date of Exam: {meta.get('date', '')}",      f"Duration: {meta.get('duration', '')}"),
        ]
        for i, (l, r) in enumerate(rows_data):
            info.rows[i].cells[0].text = l
            info.rows[i].cells[1].text = r
            for cell in info.rows[i].cells:
                if cell.paragraphs[0].runs:
                    cell.paragraphs[0].runs[0].font.size = Pt(9)
        _no_borders(info)

        _p(f"Course Code and Title: {(meta.get('course_code', '').strip() + ' – ') if meta.get('course_code', '').strip() else ''}{course_title}", bold=True, size=10)

        # CO table
        cos = meta.get("cos", [])
        if cos:
            co_tbl = doc.add_table(rows=1 + len(cos), cols=2)
            co_tbl.style = "Table Grid"
            co_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ci, txt in enumerate(["CO", "Description"]):
                cell = co_tbl.rows[0].cells[ci]
                cell.text = txt
                run = cell.paragraphs[0].runs[0]
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                _shd(cell, "4A5568")
            for i, co in enumerate(cos, 1):
                row = co_tbl.rows[i].cells
                row[0].text = co.get("id", "")
                row[1].text = co.get("text", "")
                for cell in row:
                    if cell.paragraphs[0].runs:
                        cell.paragraphs[0].runs[0].font.size = Pt(9)
                row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            co_tbl.rows[0].cells[0].width = Inches(0.6)
            co_tbl.rows[0].cells[1].width = Inches(5.6)

        leg = doc.add_paragraph(BLOOM_LEGEND)
        if leg.runs:
            leg.runs[0].font.size = Pt(7)
        leg.paragraph_format.space_after = Pt(6)

        # Questions by part
        import re as _re
        for part in paper["parts"]:
            h = doc.add_heading(f"Part {part['part']}  ({part['marks_per_question']} marks each)", level=2)
            if h.runs:
                h.runs[0].font.size = Pt(11)

            headers = ["Q.No", "Question", "CO", "BL", "Marks"]
            tbl     = doc.add_table(rows=1, cols=5)
            tbl.style = "Table Grid"
            for ci, txt in enumerate(headers):
                cell = tbl.rows[0].cells[ci]
                cell.text = txt
                run = cell.paragraphs[0].runs[0]
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                _shd(cell, "4A5568")

            for q in part["questions"]:
                raw_co = q.get("co", "")
                m = _re.match(r'(CO\d+)', str(raw_co), _re.IGNORECASE)
                co_label = m.group(1).upper() if m else raw_co
                row      = tbl.add_row().cells
                row[0].text = str(q["q_no"])
                row[1].text = q["question"]
                row[2].text = co_label
                row[3].text = BL_MAP.get(q["bloom"], q["bloom"])
                row[4].text = str(q["marks"])
                for ci, cell in enumerate(row):
                    if cell.paragraphs[0].runs:
                        cell.paragraphs[0].runs[0].font.size = Pt(9)
                    if ci != 1:
                        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph()

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    # ── Helpers ───────────────────────────────────────────────────────────────

    def _assign_bt_ranges(self, n_parts: int) -> list:
        if n_parts == 1:
            return [BT_ORDER[:]]
        if n_parts == 2:
            return [["BT1", "BT2", "BT3"], ["BT4", "BT5", "BT6"]]
        if n_parts == 3:
            return [["BT1", "BT2"], ["BT3", "BT4"], ["BT5", "BT6"]]
        chunk = math.ceil(6 / n_parts)
        return [BT_ORDER[i*chunk:min(i*chunk+chunk, 6)] or [BT_ORDER[-1]] for i in range(n_parts)]

    def _distribute(self, total: int, buckets: int) -> list:
        if buckets == 0:
            return []
        base, rem = divmod(total, buckets)
        return [base + (1 if i < rem else 0) for i in range(buckets)]

    def _generate_question(self, topic: str, subtopics: list, bt: str, marks: int, unit_title: str) -> str:
        verbs      = BT_VERBS.get(bt, "Explain")
        sub_hint   = (", ".join(subtopics[:4])) if subtopics else topic
        complexity = ("simple recall" if bt in ("BT1", "BT2") else
                      "application/problem-solving" if bt in ("BT3", "BT4") else
                      "critical evaluation or design")
        prompt = (
            f"Generate exactly ONE exam question for the topic '{topic}' (Unit: {unit_title}).\n"
            f"Bloom level: {bt} — use verbs like: {verbs}.\n"
            f"Subtopics to draw from: {sub_hint}.\n"
            f"Marks: {marks}. Complexity: {complexity}.\n"
            f"Return ONLY the question text, no numbering, no explanation."
        )
        try:
            resp = self.groq.chat.completions.create(
                messages=[{"role": "user", "content": prompt}],
                model="llama-3.3-70b-versatile",
                temperature=0.7,
                max_tokens=120,
            )
            return resp.choices[0].message.content.strip().strip('"')
        except Exception as e:
            logger.warning(f"Groq question gen failed for {topic}/{bt}: {e}")
            return f"{verbs.split('/')[0]} {topic}."

    def _build_summary(self, paper_parts: list, co_keys: list) -> dict:
        bt_count   = {bt: 0 for bt in BT_ORDER}
        co_count   = {co: 0 for co in co_keys}
        unit_count = {}
        for part in paper_parts:
            for q in part["questions"]:
                bt_count[q["bloom"]] = bt_count.get(q["bloom"], 0) + 1
                if q["co"]:
                    co_count[q["co"]] = co_count.get(q["co"], 0) + 1
                u = q["unit"]
                unit_count[u] = unit_count.get(u, 0) + 1
        return {
            "bloom_distribution": {k: v for k, v in bt_count.items() if v > 0},
            "co_distribution":    {k: v for k, v in co_count.items() if v > 0},
            "unit_distribution":  unit_count,
            "total_questions":    sum(len(p["questions"]) for p in paper_parts),
            "total_marks":        sum(q["marks"] for p in paper_parts for q in p["questions"]),
        }
