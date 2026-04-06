"""
Question paper generator — Anna University format.
Header + CO table + questions all on the SAME first page (no page break).
"""

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch, mm
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime
import tempfile

PAGE_W = A4[0]

BLOOM_LEGEND = (
    "BL – Bloom's Taxonomy Levels  "
    "(L1–Remembering, L2–Understanding, L3–Applying, "
    "L4–Analysing, L5–Evaluating, L6–Creating)"
)
BL_MAP = {"BT1": "L1", "BT2": "L2", "BT3": "L3", "BT4": "L4", "BT5": "L5", "BT6": "L6"}


def _bl(level: str) -> str:
    return BL_MAP.get(level, level or "N/A")


def _single_co(result: dict) -> str:
    """Return only the top (first) predicted CO."""
    cos = result.get("course_outcomes") or []
    return cos[0] if cos else "N/A"


class ReportGenerator:

    # ── Public API ────────────────────────────────────────────────────────────

    @staticmethod
    def generate_pdf(domain_name, results, report_type="mapping", meta=None):
        temp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
        doc  = SimpleDocTemplate(
            temp.name, pagesize=A4,
            leftMargin=20*mm, rightMargin=20*mm,
            topMargin=15*mm, bottomMargin=15*mm,
        )
        styles   = getSampleStyleSheet()
        elements = []

        # Header + CO table inline (no page break)
        if meta:
            elements += ReportGenerator._header_pdf(meta, styles)

        elements += ReportGenerator._questions_pdf(results, report_type, styles)
        doc.build(elements)
        return temp.name

    @staticmethod
    def generate_docx(domain_name, results, report_type="mapping", meta=None):
        temp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
        doc  = Document()

        for section in doc.sections:
            section.top_margin    = Inches(0.6)
            section.bottom_margin = Inches(0.6)
            section.left_margin   = Inches(0.8)
            section.right_margin  = Inches(0.8)

        # Header + CO table inline (no page break)
        if meta:
            ReportGenerator._header_docx(doc, meta)

        ReportGenerator._questions_docx(doc, results, report_type)
        doc.save(temp.name)
        return temp.name

    # ── PDF header ────────────────────────────────────────────────────────────

    @staticmethod
    def _header_pdf(meta, styles):
        usable_w = PAGE_W - 40*mm
        elems    = []

        def s(name, **kw):
            return ParagraphStyle(name, parent=styles["Normal"], **kw)

        cb = s("cb", fontSize=11, fontName="Helvetica-Bold", alignment=1, spaceAfter=1)
        cr = s("cr", fontSize=10, alignment=1, spaceAfter=1)
        sm = s("sm", fontSize=8)
        sb = s("sb", fontSize=8, fontName="Helvetica-Bold")

        elems.append(Paragraph("ANNA UNIVERSITY, CHENNAI", cb))
        elems.append(Paragraph("DEPARTMENT OF INFORMATION SCIENCE AND TECHNOLOGY", cb))
        elems.append(Spacer(1, 3))
        elems.append(Paragraph(meta.get("assessment", "Assessment Test – I"), cr))
        elems.append(Spacer(1, 6))

        prog     = meta.get("programme", "MCA (R & SS)")
        marks    = meta.get("max_marks", 50)
        date     = meta.get("date", datetime.now().strftime("%d/%m/%Y"))
        sem      = meta.get("semester", "I / III Sem")
        reg      = meta.get("regulation", "2023")
        duration = meta.get("duration", "1 hour 30 mins")

        info = Table([
            [Paragraph(f"<b>Programme:</b> {prog}", sm),   Paragraph(f"<b>Year / SEM:</b> {sem}", sm)],
            [Paragraph(f"<b>Max. Marks:</b> {marks}", sm), Paragraph(f"<b>Regulation:</b> {reg}", sm)],
            [Paragraph(f"<b>Date of Exam:</b> {date}", sm),Paragraph(f"<b>Duration:</b> {duration}", sm)],
        ], colWidths=[usable_w * 0.5, usable_w * 0.5])
        info.setStyle(TableStyle([
            ("VALIGN", (0,0),(-1,-1), "TOP"),
            ("TOPPADDING",    (0,0),(-1,-1), 2),
            ("BOTTOMPADDING", (0,0),(-1,-1), 2),
        ]))
        elems.append(info)
        elems.append(Spacer(1, 5))

        code  = meta.get("course_code", "")
        title = meta.get("course_title", "")
        elems.append(Paragraph(f"<b>Course Code and Title:</b> {code} &amp; {title}", cr))
        elems.append(Spacer(1, 6))

        cos = meta.get("cos", [])
        if cos:
            co_data = [[Paragraph("<b>CO</b>", sb), Paragraph("<b>Description</b>", sb)]]
            for co in cos:
                co_data.append([Paragraph(co.get("id",""), sm), Paragraph(co.get("text",""), sm)])
            co_tbl = Table(co_data, colWidths=[usable_w * 0.10, usable_w * 0.90])
            co_tbl.setStyle(TableStyle([
                ("GRID",          (0,0),(-1,-1), 0.5, colors.black),
                ("BACKGROUND",    (0,0),(-1, 0), colors.HexColor("#4a5568")),
                ("TEXTCOLOR",     (0,0),(-1, 0), colors.white),
                ("ALIGN",         (0,0),(0, -1), "CENTER"),
                ("VALIGN",        (0,0),(-1,-1), "TOP"),
                ("TOPPADDING",    (0,0),(-1,-1), 3),
                ("BOTTOMPADDING", (0,0),(-1,-1), 3),
            ]))
            elems.append(co_tbl)
            elems.append(Spacer(1, 4))

        elems.append(Paragraph(BLOOM_LEGEND, s("bl", fontSize=7, spaceAfter=6)))
        return elems

    # ── DOCX header ───────────────────────────────────────────────────────────

    @staticmethod
    def _header_docx(doc, meta):
        def _p(text, bold=False, size=11, align=WD_ALIGN_PARAGRAPH.CENTER):
            p = doc.add_paragraph()
            p.alignment = align
            p.paragraph_format.space_after  = Pt(1)
            p.paragraph_format.space_before = Pt(0)
            r = p.add_run(text)
            r.bold = bold
            r.font.size = Pt(size)

        _p("ANNA UNIVERSITY, CHENNAI", bold=True, size=12)
        _p("DEPARTMENT OF INFORMATION SCIENCE AND TECHNOLOGY", bold=True, size=11)
        _p(meta.get("assessment", "Assessment Test – I"), size=10)

        prog     = meta.get("programme", "MCA (R & SS)")
        marks    = meta.get("max_marks", 50)
        date     = meta.get("date", datetime.now().strftime("%d/%m/%Y"))
        sem      = meta.get("semester", "I / III Sem")
        reg      = meta.get("regulation", "2023")
        duration = meta.get("duration", "1 hour 30 mins")

        tbl = doc.add_table(rows=3, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        rows_data = [
            (f"Programme: {prog}",    f"Year / SEM: {sem}"),
            (f"Max. Marks: {marks}",  f"Regulation: {reg}"),
            (f"Date of Exam: {date}", f"Duration: {duration}"),
        ]
        for i, (l, r) in enumerate(rows_data):
            tbl.rows[i].cells[0].text = l
            tbl.rows[i].cells[1].text = r
            for cell in tbl.rows[i].cells:
                if cell.paragraphs[0].runs:
                    cell.paragraphs[0].runs[0].font.size = Pt(9)
        # Remove table borders
        for row in tbl.rows:
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for side in ("top","left","bottom","right","insideH","insideV"):
                    b = OxmlElement(f"w:{side}")
                    b.set(qn("w:val"), "none")
                    borders.append(b)
                tcPr.append(borders)

        code  = meta.get("course_code", "")
        title = meta.get("course_title", "")
        _p(f"Course Code and Title: {code} & {title}", bold=True, size=10)

        cos = meta.get("cos", [])
        if cos:
            co_tbl = doc.add_table(rows=1 + len(cos), cols=2)
            co_tbl.style = "Table Grid"
            co_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

            hdr = co_tbl.rows[0].cells
            for ci, txt in enumerate(["CO", "Description"]):
                hdr[ci].text = txt
                run = hdr[ci].paragraphs[0].runs[0]
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                hdr[ci].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "4A5568")
                hdr[ci]._tc.get_or_add_tcPr().append(shd)

            for i, co in enumerate(cos, 1):
                row = co_tbl.rows[i].cells
                row[0].text = co.get("id", "")
                row[1].text = co.get("text", "")
                for cell in row:
                    if cell.paragraphs[0].runs:
                        cell.paragraphs[0].runs[0].font.size = Pt(9)
                row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for row in co_tbl.rows:
                row.cells[0].width = Inches(0.6)
                row.cells[1].width = Inches(5.6)

        leg = doc.add_paragraph(BLOOM_LEGEND)
        if leg.runs:
            leg.runs[0].font.size = Pt(7)
        leg.paragraph_format.space_after = Pt(6)

    # ── PDF questions ─────────────────────────────────────────────────────────

    @staticmethod
    def _questions_pdf(results, report_type, styles):
        usable_w = PAGE_W - 40*mm
        elems    = []

        cs = ParagraphStyle("cs", parent=styles["Normal"], fontSize=9, leading=11, wordWrap="CJK")
        ps = ParagraphStyle("ps", parent=styles["Heading2"], fontSize=11, spaceAfter=6, spaceBefore=8)

        parts = ReportGenerator.detect_parts(results)

        for part_name in ["Part A", "Part B", "Part C"]:
            if part_name not in parts:
                continue

            elems.append(Paragraph(part_name, ps))

            if report_type == "mapping":
                headers    = ["Q.No", "CO", "BL"]
                col_widths = [usable_w*0.12, usable_w*0.55, usable_w*0.33]
            else:
                headers    = ["Q.No", "Question", "CO", "BL"]
                col_widths = [usable_w*0.07, usable_w*0.63, usable_w*0.16, usable_w*0.14]

            rows = [[Paragraph(f"<b>{h}</b>", cs) for h in headers]]
            for r in parts[part_name]:
                co  = _single_co(r)
                bl  = _bl(r.get("bloom_level", ""))
                if report_type == "mapping":
                    rows.append([
                        Paragraph(str(r.get("question_number", "-")), cs),
                        Paragraph(co, cs),
                        Paragraph(bl, cs),
                    ])
                else:
                    rows.append([
                        Paragraph(str(r.get("question_number", "-")), cs),
                        Paragraph(r.get("question", ""), cs),
                        Paragraph(co, cs),
                        Paragraph(bl, cs),
                    ])

            tbl = Table(rows, colWidths=col_widths)
            tbl.setStyle(TableStyle([
                ("BACKGROUND",    (0,0),(-1, 0), colors.HexColor("#4a5568")),
                ("TEXTCOLOR",     (0,0),(-1, 0), colors.white),
                ("FONTNAME",      (0,0),(-1, 0), "Helvetica-Bold"),
                ("FONTSIZE",      (0,0),(-1, 0), 9),
                ("ALIGN",         (0,0),(-1,-1), "CENTER"),
                ("VALIGN",        (0,0),(-1,-1), "MIDDLE"),
                ("TOPPADDING",    (0,0),(-1,-1), 4),
                ("BOTTOMPADDING", (0,0),(-1,-1), 4),
                ("GRID",          (0,0),(-1,-1), 0.5, colors.grey),
                ("ROWBACKGROUNDS",(0,1),(-1,-1), [colors.white, colors.HexColor("#f7fafc")]),
            ]))
            elems.append(tbl)
            elems.append(Spacer(1, 0.2*inch))

        return elems

    # ── DOCX questions ────────────────────────────────────────────────────────

    @staticmethod
    def _questions_docx(doc, results, report_type):
        parts = ReportGenerator.detect_parts(results)

        for part_name in ["Part A", "Part B", "Part C"]:
            if part_name not in parts:
                continue

            h = doc.add_heading(part_name, level=2)
            if h.runs:
                h.runs[0].font.size = Pt(11)

            headers = ["Q.No", "CO", "BL"] if report_type == "mapping" else ["Q.No", "Question", "CO", "BL"]
            tbl = doc.add_table(rows=1, cols=len(headers))
            tbl.style = "Table Grid"

            for i, txt in enumerate(headers):
                cell = tbl.rows[0].cells[i]
                cell.text = txt
                run = cell.paragraphs[0].runs[0]
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "4A5568")
                cell._tc.get_or_add_tcPr().append(shd)

            for r in parts[part_name]:
                co  = _single_co(r)
                bl  = _bl(r.get("bloom_level", ""))
                row = tbl.add_row().cells
                if report_type == "mapping":
                    row[0].text = str(r.get("question_number", "-"))
                    row[1].text = co
                    row[2].text = bl
                else:
                    row[0].text = str(r.get("question_number", "-"))
                    row[1].text = r.get("question", "")
                    row[2].text = co
                    row[3].text = bl
                for cell in row:
                    if cell.paragraphs[0].runs:
                        cell.paragraphs[0].runs[0].font.size = Pt(9)

            doc.add_paragraph()

    # ── Helpers ───────────────────────────────────────────────────────────────

    @staticmethod
    def detect_parts(results):
        parts = {}
        for r in results:
            parts.setdefault(r.get("part", "Part A"), []).append(r)
        return parts
